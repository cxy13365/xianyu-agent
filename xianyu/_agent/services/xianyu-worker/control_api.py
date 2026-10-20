from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import time
from collections import deque
from contextlib import asynccontextmanager
from datetime import datetime
from typing import Any, Literal

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from loguru import logger
from pydantic import BaseModel, Field

from config_store import ConfigStore
from main import XianyuLive
from context_manager import ChatContextManager
from pipeline import AgentPipeline, KnowledgeBase, PipelineReplyBot
from faq_tools import search_faq
from graph_persistence import GraphPersistence
from utils.xianyu_utils import trans_cookies


class ConfigPayload(BaseModel):
    cookie: str | None = None
    api_key: str | None = None
    model_base_url: str | None = None
    model_name: str | None = None
    agent_mode: Literal["legacy", "shadow", "langgraph"] | None = None


class PipelinePayload(BaseModel):
    message: str = Field(min_length=1, max_length=1000)
    item_id: str = "*"
    item_title: str = "测试商品"
    item_desc: str = ""
    context: list[dict[str, Any]] = []


class ReplyPolicyPayload(BaseModel):
    simulate_human_typing: bool | None = None
    low_confidence_handoff: bool | None = None
    confidence_threshold: float | None = Field(default=None, ge=0, le=1)
    block_off_platform: bool | None = None
    night_auto_reply: bool | None = None


class ManualReplyPayload(BaseModel):
    chat_id: str = Field(min_length=1, max_length=200)
    to_user_id: str = Field(min_length=1, max_length=200)
    item_id: str = "*"
    message: str = Field(min_length=1, max_length=2000)


class ApprovalPayload(BaseModel):
    reply: str = Field(default="", max_length=2000)


class FaqPayload(BaseModel):
    scene: str = "通用"
    question: str = Field(min_length=1, max_length=1000)
    answer: str = Field(min_length=1, max_length=3000)
    keywords: str = ""
    enabled: bool = True


class FaqSearchPayload(BaseModel):
    query: str = Field(min_length=1, max_length=1000)
    limit: int = Field(default=3, ge=1, le=5, strict=True)


class Runtime:
    def __init__(self):
        self.store = ConfigStore()
        self.kb = KnowledgeBase()
        self.graph_store = GraphPersistence()
        self.events: deque[dict] = deque(maxlen=300)
        self.live: XianyuLive | None = None
        self.task: asyncio.Task | None = None
        self.status = "disconnected"
        self.last_error = ""
        self.connected_at: float | None = None

    def emit(self, event: dict):
        event.setdefault("at", time.time())
        self.events.appendleft(event)

    async def connect(self):
        if self.task and not self.task.done():
            return
        config = self.store.load()
        cookie = str(config.get("cookie", ""))
        parsed = trans_cookies(cookie)
        if not cookie or "unb" not in parsed:
            raise HTTPException(400, "Cookie 缺少 unb 字段，请从已登录的 goofish.com 重新复制完整 Cookie")
        if not config.get("api_key"):
            raise HTTPException(400, "请先配置模型 API Key")
        ConfigStore.apply_environment(config)
        self.status, self.last_error = "connecting", ""
        try:
            bot = PipelineReplyBot(
                event_sink=self.emit,
                mode=str(config.get("agent_mode") or "shadow"),
                persistence=self.graph_store,
                kb=self.kb,
            )
            self.live = XianyuLive(cookie, bot=bot)
        except Exception as exc:
            self.status, self.last_error = "error", str(exc)
            raise HTTPException(400, f"初始化失败：{exc}") from exc

        async def runner():
            try:
                self.status = "connected"
                self.connected_at = time.time()
                self.emit({"type": "connector.connected", "message": "闲鱼消息监听已启动"})
                await self.live.main()
            except asyncio.CancelledError:
                raise
            except BaseException as exc:
                self.status, self.last_error = "error", str(exc)
                self.emit({"type": "connector.error", "message": str(exc)})
                logger.exception("闲鱼连接任务退出")
            finally:
                if self.status != "error":
                    self.status = "disconnected"

        self.task = asyncio.create_task(runner(), name="xianyu-live")

    async def disconnect(self):
        if self.live and self.live.ws:
            try:
                await self.live.ws.close()
            except Exception:
                pass
        if self.task and not self.task.done():
            self.task.cancel()
            try:
                await self.task
            except asyncio.CancelledError:
                pass
        self.task, self.live = None, None
        self.status = "disconnected"
        self.emit({"type": "connector.disconnected", "message": "用户停止自动回复"})


runtime = Runtime()


def _pending_handoffs() -> list[dict]:
    """汇总当前连接中的人工接管和待批准 Graph 会话。"""
    if runtime.status != "connected" or not runtime.live:
        return []
    conversations = {
        item["chat_id"]: item
        for item in runtime.live.context_manager.list_conversations(200)
    }
    pending: dict[str, dict] = {}
    for chat_id in list(runtime.live.manual_mode_conversations):
        if not runtime.live.is_manual_mode(chat_id):
            continue
        conversation = conversations.get(chat_id, {})
        reason = "高风险用户自动转人工" if conversation.get("risk_level") == "high" else "人工接管"
        pending[chat_id] = {
            "chat_id": chat_id,
            "message": str(conversation.get("last_message") or "等待人工回复"),
            "reason": reason,
            "updated_at": str(conversation.get("timestamp") or ""),
        }
    for run in runtime.graph_store.list_active_runs():
        conversation = conversations.get(run["chat_id"], {})
        reason = {
            "interrupted": "等待批准回复",
            "paused": "Graph 已暂停",
            "awaiting_approval": "等待批准回复",
        }.get(run["status"], "需要人工处理")
        pending[run["chat_id"]] = {
            "chat_id": run["chat_id"],
            "message": str(conversation.get("last_message") or "等待人工处理"),
            "reason": reason,
            "updated_at": run["updated_at"],
        }
    def updated_timestamp(value: object) -> float:
        if isinstance(value, (int, float)):
            return float(value)
        try:
            return datetime.fromisoformat(str(value)).timestamp()
        except ValueError:
            return 0.0

    return sorted(pending.values(), key=lambda item: updated_timestamp(item["updated_at"]), reverse=True)


@asynccontextmanager
async def lifespan(_: FastAPI):
    config = runtime.store.load()
    ConfigStore.apply_environment(config)
    yield
    await runtime.disconnect()


app = FastAPI(title="Xianyu Agent Desk Worker", version="1.0.0", lifespan=lifespan)
origins = [x.strip() for x in os.getenv("ALLOWED_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000").split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=False,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["*"],
)


@app.get("/health")
def health():
    return {"ok": True, "service": "xianyu-agent-worker", "status": runtime.status}


@app.get("/api/status")
def status():
    return {
        "status": runtime.status,
        "connected": runtime.status == "connected",
        "connected_at": runtime.connected_at,
        "last_error": runtime.last_error,
        "account_id": runtime.live.myid[-4:] if runtime.live else None,
        "agent_mode": runtime.live.bot.mode if runtime.live else str(runtime.store.load().get("agent_mode") or "shadow"),
    }


@app.get("/api/config")
def get_config():
    return ConfigStore.masked(runtime.store.load())


@app.put("/api/config")
def put_config(payload: ConfigPayload):
    values = payload.model_dump(exclude_none=True)
    saved = runtime.store.save(values)
    ConfigStore.apply_environment(saved)
    if runtime.live:
        runtime.live.bot.reload_prompts()
        runtime.live.bot.set_mode(str(saved.get("agent_mode") or "shadow"))
    return ConfigStore.masked(saved)


@app.get("/api/agent-mode")
def get_agent_mode():
    mode = runtime.live.bot.mode if runtime.live else str(runtime.store.load().get("agent_mode") or "shadow")
    return {"mode": mode, "available": ["legacy", "shadow", "langgraph"]}


@app.put("/api/agent-mode")
def put_agent_mode(payload: ConfigPayload):
    if payload.agent_mode is None:
        raise HTTPException(422, "缺少 agent_mode")
    saved = runtime.store.save({"agent_mode": payload.agent_mode})
    if runtime.live:
        runtime.live.bot.set_mode(payload.agent_mode)
    runtime.emit({"type": "agent.mode_changed", "mode": payload.agent_mode})
    return {"mode": saved["agent_mode"], "available": ["legacy", "shadow", "langgraph"]}


@app.post("/api/connect")
async def connect():
    await runtime.connect()
    return {"ok": True, "status": runtime.status, "message": "连接任务已启动"}


@app.post("/api/disconnect")
async def disconnect():
    await runtime.disconnect()
    return {"ok": True, "status": runtime.status}


@app.post("/api/pipeline/test")
async def pipeline_test(payload: PipelinePayload):
    ConfigStore.apply_environment(runtime.store.load())
    pipeline = AgentPipeline(kb=runtime.kb, event_sink=runtime.emit)
    item = f"商品：{payload.item_title}\n描述：{payload.item_desc}"
    return await asyncio.to_thread(
        pipeline.process,
        payload.message,
        item,
        payload.context,
        item_id=payload.item_id,
        chat_id="console-test",
    )


@app.get("/api/reply-policy")
def get_reply_policy():
    return runtime.kb.get_reply_policy()


@app.put("/api/reply-policy")
def put_reply_policy(payload: ReplyPolicyPayload):
    policy = runtime.kb.set_reply_policy(**payload.model_dump())
    if runtime.live:
        runtime.live.simulate_human_typing = policy["simulate_human_typing"]
    runtime.emit({"type": "reply_policy.updated", "policy": policy})
    return policy


@app.get("/api/conversations")
def conversations(limit: int = 50):
    if not runtime.live:
        return {"items": []}
    items = runtime.live.context_manager.list_conversations(limit)
    for item in items:
        item["mode"] = "manual" if runtime.live.is_manual_mode(item["chat_id"]) else "auto"
        graph_run = runtime.graph_store.latest_run(item["chat_id"])
        if graph_run:
            item.update({
                "thread_id": graph_run["thread_id"], "graph_status": graph_run["status"],
                "graph_nodes": graph_run["nodes"], "pending_reply": graph_run["draft_reply"],
            })
    return {"items": items}


@app.get("/api/conversations/{chat_id}/messages")
def conversation_messages(chat_id: str):
    if not runtime.live:
        return {"items": []}
    return {"items": runtime.live.context_manager.get_recent_messages_by_chat(chat_id, limit=20)}


@app.get("/api/graph/threads/{chat_id}")
def graph_thread(chat_id: str):
    run = runtime.graph_store.latest_run(chat_id)
    if not run:
        raise HTTPException(404, "该会话还没有 Graph 运行记录")
    return run


@app.post("/api/graph/threads/{chat_id}/pause")
def pause_graph(chat_id: str):
    if not runtime.live:
        raise HTTPException(409, "闲鱼 Worker 尚未连接")
    try:
        run = runtime.live.bot.pause(chat_id)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(409, str(exc)) from exc
    runtime.emit({"type": "graph.paused", "chat_id": chat_id, "thread_id": run["thread_id"]})
    return run


@app.post("/api/graph/threads/{chat_id}/resume")
def resume_graph(chat_id: str):
    if not runtime.live:
        raise HTTPException(409, "闲鱼 Worker 尚未连接")
    try:
        result = runtime.live.bot.resume(chat_id)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(409, str(exc)) from exc
    runtime.emit({"type": "graph.resumed", "chat_id": chat_id, "result": result})
    return result


async def _approve_graph(chat_id: str, modified_reply: str = ""):
    if not runtime.live or not runtime.live.ws or runtime.status != "connected":
        raise HTTPException(409, "闲鱼 Worker 尚未连接")
    run = runtime.graph_store.latest_run(chat_id)
    if not run or not run.get("to_user_id"):
        raise HTTPException(409, "待批准任务缺少买家发送目标")
    try:
        result = runtime.live.bot.approve(chat_id, modified_reply)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(409, str(exc)) from exc
    reply = str(result.get("reply") or "").strip()
    if not reply or reply == "-":
        raise HTTPException(409, "Graph 没有生成可发送回复")
    await runtime.live.send_msg(runtime.live.ws, chat_id, run["to_user_id"], reply)
    runtime.live.context_manager.add_message_by_chat(
        chat_id, runtime.live.myid, run["item_id"], "assistant", reply,
    )
    runtime.emit({
        "type": "graph.reply_approved", "chat_id": chat_id,
        "thread_id": run["thread_id"], "modified": bool(modified_reply), "reply": reply,
    })
    return {"ok": True, "chat_id": chat_id, "reply": reply, "result": result}


@app.post("/api/graph/threads/{chat_id}/approve")
async def approve_graph(chat_id: str):
    return await _approve_graph(chat_id)


@app.post("/api/graph/threads/{chat_id}/modify")
async def modify_graph(chat_id: str, payload: ApprovalPayload):
    if not payload.reply.strip():
        raise HTTPException(422, "修改后的回复不能为空")
    return await _approve_graph(chat_id, payload.reply.strip())


@app.get("/api/shadow-comparisons")
def shadow_comparisons(limit: int = 50):
    return {"items": runtime.graph_store.list_shadow(limit)}


@app.post("/api/conversations/{chat_id}/takeover")
def takeover(chat_id: str):
    if not runtime.live:
        raise HTTPException(409, "闲鱼 Worker 尚未连接")
    runtime.live.enter_manual_mode(chat_id)
    runtime.emit({"type": "conversation.takeover", "chat_id": chat_id})
    return {"ok": True, "chat_id": chat_id, "mode": "manual"}


@app.post("/api/conversations/{chat_id}/release")
def release(chat_id: str):
    if not runtime.live:
        raise HTTPException(409, "闲鱼 Worker 尚未连接")
    runtime.live.exit_manual_mode(chat_id)
    runtime.emit({"type": "conversation.release", "chat_id": chat_id})
    return {"ok": True, "chat_id": chat_id, "mode": "auto"}


@app.post("/api/conversations/reply")
async def manual_reply(payload: ManualReplyPayload):
    if not runtime.live or not runtime.live.ws or runtime.status != "connected":
        raise HTTPException(409, "闲鱼 Worker 尚未连接")
    if not runtime.live.is_manual_mode(payload.chat_id):
        raise HTTPException(409, "请先接管该会话再发送人工回复")
    await runtime.live.send_msg(runtime.live.ws, payload.chat_id, payload.to_user_id, payload.message)
    runtime.live.context_manager.add_message_by_chat(
        payload.chat_id, runtime.live.myid, payload.item_id, "assistant", payload.message
    )
    runtime.emit({"type": "conversation.manual_reply", "chat_id": payload.chat_id, "message": payload.message})
    return {"ok": True, "chat_id": payload.chat_id}


@app.get("/api/faqs")
def list_faqs():
    return {"items": runtime.kb.list_faqs()}


@app.get("/api/faq-suggestions")
def list_faq_suggestions(limit: int = 100):
    return {"items": runtime.kb.list_question_suggestions("pending", limit)}


@app.post("/api/faq-suggestions/{suggestion_id}/ignore")
def ignore_faq_suggestion(suggestion_id: int):
    suggestion = runtime.kb.set_question_suggestion_status(suggestion_id, "ignored")
    if suggestion is None:
        raise HTTPException(404, "自动推送问题不存在")
    runtime.emit({"type": "faq.suggestion_ignored", "suggestion_id": suggestion_id})
    return suggestion


@app.post("/api/faq-suggestions/{suggestion_id}/resolve")
def resolve_faq_suggestion(suggestion_id: int):
    suggestion = runtime.kb.set_question_suggestion_status(suggestion_id, "resolved")
    if suggestion is None:
        raise HTTPException(404, "自动推送问题不存在")
    runtime.emit({"type": "faq.suggestion_resolved", "suggestion_id": suggestion_id})
    return suggestion


@app.post("/api/faqs/search")
def search_faqs(payload: FaqSearchPayload):
    if not payload.query.strip():
        raise HTTPException(422, "请输入非空的买家问题")
    ConfigStore.apply_environment(runtime.store.load())
    pipeline = AgentPipeline(kb=runtime.kb)
    return search_faq(pipeline, payload.query, payload.limit)


@app.post("/api/faqs")
def add_faq(payload: FaqPayload):
    faq_id = runtime.kb.add_faq(
        payload.scene, payload.question, payload.answer, payload.keywords, payload.enabled
    )
    return {"id": faq_id, **payload.model_dump()}


@app.put("/api/faqs/{faq_id}")
def update_faq(faq_id: int, payload: FaqPayload):
    faq = runtime.kb.update_faq(
        faq_id,
        payload.scene,
        payload.question,
        payload.answer,
        payload.keywords,
        payload.enabled,
    )
    if faq is None:
        raise HTTPException(404, "问答不存在")
    return faq


@app.delete("/api/faqs/{faq_id}")
def delete_faq(faq_id: int):
    if not runtime.kb.delete_faq(faq_id):
        raise HTTPException(404, "问答不存在")
    runtime.emit({"type": "faq.deleted", "faq_id": faq_id})
    return {"ok": True, "id": faq_id}


def _rows_from_upload(filename: str, content: bytes) -> list[dict]:
    lower = filename.lower()
    if lower.endswith(".json"):
        data = json.loads(content.decode("utf-8-sig"))
        return data if isinstance(data, list) else data.get("items", [])
    if lower.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    if lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(500, "缺少 openpyxl，无法读取 XLSX") from exc
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = book.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    raise HTTPException(400, "仅支持 CSV、JSON、XLSX 文件")


@app.post("/api/faqs/import")
async def import_faqs(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(400, "文件名为空")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "文件不能超过 10MB")
    try:
        rows = _rows_from_upload(file.filename, content)
    except (ValueError, json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise HTTPException(400, f"文件解析失败：{exc}") from exc
    imported, errors = 0, []
    for index, row in enumerate(rows, start=2):
        question = str(row.get("question") or row.get("问题") or "").strip()
        answer = str(row.get("answer") or row.get("答案") or "").strip()
        if not question or not answer:
            errors.append({"row": index, "error": "缺少 question/answer"})
            continue
        runtime.kb.add_faq(str(row.get("scene") or row.get("场景") or "通用"), question, answer, str(row.get("keywords") or row.get("关键词") or ""))
        imported += 1
    return {"imported": imported, "errors": errors[:50]}


@app.get("/api/traces")
def traces(limit: int = 50):
    return {"items": runtime.kb.recent_traces(max(1, min(limit, 200)))}


@app.get("/api/events")
def events(limit: int = 100):
    return {"items": list(runtime.events)[: max(1, min(limit, 300))]}


@app.get("/api/metrics")
def metrics():
    today = datetime.now().date().isoformat()
    context_manager = runtime.live.context_manager if runtime.live else ChatContextManager()
    activity = context_manager.daily_activity_metrics(today)
    pipeline_metrics = runtime.kb.daily_pipeline_metrics(today)
    faq_metrics = runtime.kb.daily_faq_metrics(today)
    agreement = runtime.graph_store.daily_intent_agreement(today)
    pending_handoffs = _pending_handoffs()
    buyer_conversation_count = activity["buyer_conversation_count"]
    takeover_count = activity["takeover_conversation_count"]
    return {
        "date": today,
        "reply_count": activity["reply_count"],
        "average_response_ms": pipeline_metrics["average_response_ms"],
        "response_latency_sample_count": pipeline_metrics["response_latency_sample_count"],
        **faq_metrics,
        "faq_enabled_count": sum(1 for faq in runtime.kb.list_faqs() if faq["enabled"]),
        "intent_agreement_rate": agreement["intent_agreement_rate"],
        "intent_agreement_sample_count": agreement["intent_agreement_sample_count"],
        "takeover_rate": round(100 * takeover_count / buyer_conversation_count, 1) if buyer_conversation_count else None,
        "takeover_conversation_count": takeover_count,
        "buyer_conversation_count": buyer_conversation_count,
        "takeover_tracking_started_at": activity["takeover_tracking_started_at"],
        "connected": runtime.status == "connected",
        "faq_count": len(runtime.kb.list_faqs()),
        "pending_handoff_count": len(pending_handoffs),
        "pending_handoffs": pending_handoffs[:3],
        "last_error": runtime.last_error,
    }


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("control_api:app", host=os.getenv("HOST", "127.0.0.1"), port=int(os.getenv("PORT", "8787")), reload=False)
